#!/usr/bin/env python3
"""Read mutual-fund NAV data from SEC and print MTF reports."""

from __future__ import annotations

import json
import math
import os
import re
import sys
import tempfile
import time
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Any, Optional
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode, urlparse
from urllib.request import Request, urlopen
from zipfile import BadZipFile

import openpyxl


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ENV_FILE = os.path.join(SCRIPT_DIR, ".env")
DOWNLOAD_DIR = os.path.join(SCRIPT_DIR, "downloads")
WORKBOOK_PATH = os.path.join(DOWNLOAD_DIR, "All-ThaiMutualFunds-MTFS.xlsx")
MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
HISTORY_DAYS = 1825
MAX_NAV_PAGES = 60
MAX_PROFILE_PAGES = 300
MAX_FACTSHEET_PAGES = 300
SEC_BASE_URL = "https://api.sec.or.th/v2/fund"
FINNOMENA_URL = "https://www.finnomena.com/fund/{}/graph"
BANGKOK_TIMEZONE = timezone(timedelta(hours=7))
CLASS_FIELDS = (
    "fund_class_name", "fundClassName", "class_abbr_name", "classAbbrName",
    "class_name", "className", "unit_type", "unitType", "cls_abbr_name",
    "clsAbbrName", "fund_class", "fundClass", "proj_id_class", "class_id", "classId",
)
NON_STANDARD_CLASS_SUFFIXES = ("USD", "SSF", "RMF", "UH", "UHA", "UHI", "P", "(E)")
TIMEFRAMES = {
    "W": (30 * 86400, 0),
    "D": (7 * 86400, 3 * 86400),
    "4H": (86400, 0),
    "1H": (86400, 0),
}
TF_CONFIG = {
    "W": dict(rsi=14, trend=4, fast=9, slow=21, macd_fast=12, macd_slow=26,
              macd_signal=9, buy_rsi=35, sell_rsi=75, vol=4, rs_lookback=4),
    "D": dict(rsi=14, trend=5, fast=9, slow=21, macd_fast=12, macd_slow=26,
              macd_signal=9, buy_rsi=38, sell_rsi=72, vol=7, rs_lookback=7),
    "4H": dict(rsi=14, trend=10, fast=9, slow=21, macd_fast=12, macd_slow=26,
               macd_signal=9, buy_rsi=40, sell_rsi=70, vol=20, rs_lookback=20),
    "1H": dict(rsi=7, trend=5, fast=5, slow=10, macd_fast=6, macd_slow=13,
               macd_signal=5, buy_rsi=45, sell_rsi=65, vol=8, rs_lookback=8),
}
FUND_CODE_RE = re.compile(
    r"(?=.{1,40}\Z)[A-Za-z0-9](?:[A-Za-z0-9._&()/#-]| [A-Za-z0-9._&()/#-])*\Z"
)
SHEET_ID_RE = re.compile(r"/spreadsheets/d/([A-Za-z0-9_-]+)(?:/|$)")


class FundDataError(Exception):
    """An expected input or data-provider error suitable for CLI output."""


@dataclass(frozen=True)
class FundRow:
    code: str
    project_id: str
    class_hint: str
    row_number: int


@dataclass
class NavSeries:
    timestamps: list[int]
    closes: list[float]
    units: list[float]
    class_warning: str = ""
    class_name: str = ""


def load_config() -> dict[str, str]:
    config = {key: os.environ[key] for key in ("SET_DATA_URL", "SET_DATA_SHEET", "SEC_API_KEY_FUND")
              if os.environ.get(key)}
    try:
        with open(ENV_FILE, encoding="utf-8") as env_file:
            for raw_line in env_file:
                line = raw_line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                key, value = line.split("=", 1)
                key, value = key.strip(), value.strip()
                if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
                    value = value[1:-1]
                if key in ("SET_DATA_URL", "SET_DATA_SHEET", "SEC_API_KEY_FUND"):
                    config.setdefault(key, value)
    except FileNotFoundError as exc:
        if not config.get("SET_DATA_URL"):
            raise FundDataError("ไม่พบไฟล์ .env; กรุณาคัดลอก .env.example เป็น .env") from exc
    except OSError as exc:
        raise FundDataError(f"อ่านไฟล์ .env ไม่สำเร็จ: {exc}") from exc

    missing = [key for key in ("SET_DATA_URL", "SEC_API_KEY_FUND") if not config.get(key, "").strip()]
    if missing:
        raise FundDataError("กรุณากำหนดค่าใน .env: " + ", ".join(missing))
    config.setdefault("SET_DATA_SHEET", "All-ThaiMutualFunds-MTFS")
    return config


def spreadsheet_export_url(sheet_url: str) -> str:
    parsed = urlparse(sheet_url)
    match = SHEET_ID_RE.search(parsed.path)
    if parsed.scheme != "https" or parsed.hostname != "docs.google.com" or not match:
        raise FundDataError("SET_DATA_URL ต้องเป็น URL ของ Google Sheets ที่เข้าถึงได้ด้วยลิงก์")
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=xlsx"


def _read_url(url: str, headers: dict[str, str], timeout: int = 30) -> bytes:
    request = Request(url, headers=headers)
    try:
        with urlopen(request, timeout=timeout) as response:
            return response.read(MAX_WORKBOOK_BYTES + 1)
    except HTTPError as exc:
        raise FundDataError(f"ดาวน์โหลด Google Sheets ไม่สำเร็จ (HTTP {exc.code})") from exc
    except URLError as exc:
        raise FundDataError(f"เชื่อมต่อ Google Sheets ไม่สำเร็จ: {exc.reason}") from exc
    except TimeoutError as exc:
        raise FundDataError("หมดเวลาขณะดาวน์โหลด Google Sheets") from exc
    except OSError as exc:
        raise FundDataError(f"อ่าน workbook จาก Google Sheets ไม่สำเร็จ: {exc}") from exc


def _save_workbook_snapshot(content: bytes) -> None:
    temporary_path = None
    try:
        os.makedirs(DOWNLOAD_DIR, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            mode="wb",
            dir=DOWNLOAD_DIR,
            prefix=".All-ThaiMutualFunds-MTFS-",
            suffix=".tmp",
            delete=False,
        ) as temporary_file:
            temporary_path = temporary_file.name
            temporary_file.write(content)
        os.replace(temporary_path, WORKBOOK_PATH)
    except OSError as exc:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)
        raise FundDataError(f"บันทึก snapshot ลง downloads ไม่สำเร็จ: {exc}") from exc


def load_fund_rows(config: dict[str, str], selected: Optional[list[str]]) -> list[FundRow]:
    url = spreadsheet_export_url(config["SET_DATA_URL"])
    content = _read_url(url, {
        "User-Agent": "TFOSvMF/mutual-fund-cli",
        "Cache-Control": "no-cache, no-store, max-age=0",
        "Pragma": "no-cache",
    })
    if len(content) > MAX_WORKBOOK_BYTES:
        raise FundDataError(f"ไฟล์ workbook มีขนาดเกิน {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB")
    _save_workbook_snapshot(content)
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, ValueError, OSError) as exc:
        raise FundDataError("ไฟล์ที่ดาวน์โหลดไม่ใช่ workbook XLSX ที่อ่านได้") from exc
    sheet_name = config["SET_DATA_SHEET"]
    if sheet_name not in workbook.sheetnames:
        raise FundDataError(f"ไม่พบชีต {sheet_name!r} ใน workbook")
    sheet = workbook[sheet_name]
    header_row = None
    indexes: dict[str, int] = {}
    for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        headers = [str(value).strip().upper() if value is not None else "" for value in values]
        if "AKA" in headers:
            header_row = row_number
            indexes = {header: index for index, header in enumerate(headers) if header}
            break
    if header_row is None:
        raise FundDataError("ไม่พบหัวคอลัมน์ AKA ใน 10 แถวแรก")
    if "PROJECT ID" not in indexes:
        raise FundDataError("ขาดคอลัมน์ PROJECT ID ในชีต")

    requested = {symbol.strip().upper() for symbol in selected} if selected else None
    rows: list[FundRow] = []
    seen: set[str] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True),
                                         header_row + 1):
        raw_code = values[indexes["AKA"]] if indexes["AKA"] < len(values) else None
        code = str(raw_code).strip() if raw_code is not None else ""
        if not code:
            continue
        if code.upper().endswith("(FUND)"):
            code = code[:-6].strip()
        if not FUND_CODE_RE.fullmatch(code):
            raise FundDataError(f"แถว {row_number}: AKA มีรูปแบบไม่ถูกต้อง")
        normalized = code.upper()
        if normalized in seen:
            raise FundDataError(f"แถว {row_number}: AKA ซ้ำ ({code})")
        seen.add(normalized)
        if requested is not None and normalized not in requested:
            continue
        raw_project = values[indexes["PROJECT ID"]] if indexes["PROJECT ID"] < len(values) else None
        if isinstance(raw_project, float) and raw_project.is_integer():
            raw_project = int(raw_project)
        parts = str(raw_project or "").strip().split("|", 1)
        project_id = parts[0].strip()
        if project_id.upper() in ("PROJECT ID", "PROJ_ID"):
            project_id = ""
        class_hint = parts[1].strip() if len(parts) > 1 else ""
        rows.append(FundRow(code, project_id, class_hint, row_number))
    workbook.close()
    if requested:
        found = {row.code.upper() for row in rows}
        missing = sorted(requested - found)
        if missing:
            raise FundDataError("ไม่พบ AKA ในชีต: " + ", ".join(missing))
    if not rows:
        if requested:
            raise FundDataError("ไม่พบ AKA ที่ระบุในชีต: " + ", ".join(selected or []))
        raise FundDataError("ไม่พบแถวกองทุนที่มี AKA ในชีต")
    return rows


class SecClient:
    def __init__(self, api_key: str):
        self._headers = {
            "Ocp-Apim-Subscription-Key": api_key,
            "User-Agent": "TFOSvMF/mutual-fund-cli",
        }

    def _json(self, url: str) -> tuple[int, Optional[dict[str, Any]]]:
        request = Request(url, headers=self._headers)
        try:
            with urlopen(request, timeout=30) as response:
                status = response.status
                body = response.read(10 * 1024 * 1024 + 1)
        except HTTPError as exc:
            status, body = exc.code, exc.read(10 * 1024 * 1024 + 1)
        except URLError as exc:
            raise FundDataError(f"เชื่อมต่อ SEC Open API ไม่สำเร็จ: {exc.reason}") from exc
        except TimeoutError as exc:
            raise FundDataError("หมดเวลารอ SEC Open API") from exc
        except OSError as exc:
            raise FundDataError(f"อ่านคำตอบจาก SEC Open API ไม่สำเร็จ: {exc}") from exc
        if len(body) > 10 * 1024 * 1024:
            raise FundDataError("คำตอบจาก SEC Open API มีขนาดเกิน 10 MB")
        if status == 204 or not body:
            return status, None
        if status != 200:
            raise FundDataError(f"SEC Open API ตอบกลับ HTTP {status}")
        try:
            payload = json.loads(body.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise FundDataError("SEC Open API ส่งข้อมูล JSON ไม่ถูกต้อง") from exc
        if not isinstance(payload, dict):
            raise FundDataError("รูปแบบคำตอบจาก SEC Open API ไม่ถูกต้อง")
        return status, payload

    def profiles(self) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        cursor: Optional[str] = None
        for page in range(MAX_PROFILE_PAGES):
            url = f"{SEC_BASE_URL}/general-info/profiles"
            if cursor:
                url += "?" + urlencode({"next_cursor": cursor})
            _, payload = self._json(url)
            if payload is None:
                raise FundDataError("SEC Open API ไม่ส่งรายการ fund profiles")
            page_items = payload.get("items", [])
            if not isinstance(page_items, list):
                raise FundDataError("รูปแบบรายการ fund profiles จาก SEC ไม่ถูกต้อง")
            items.extend(item for item in page_items if isinstance(item, dict))
            cursor = payload.get("next_cursor") or None
            if not cursor:
                return items
            time.sleep(0.3)
        raise FundDataError(f"SEC fund profiles เกินขีดจำกัด {MAX_PROFILE_PAGES} หน้า")

    def find_project_id(self, profiles: list[dict[str, Any]], fund_code: str) -> str:
        def name(profile: dict[str, Any]) -> str:
            for key in ("proj_abbr_name", "projAbbrName", "fund_abbr_name", "fundAbbrName",
                        "abbr_name", "abbrName", "proj_name_en"):
                value = str(profile.get(key) or "").strip()
                if value:
                    return value.upper()
            return ""

        def project_id(profile: dict[str, Any]) -> str:
            return str(profile.get("proj_id") or profile.get("projId") or
                       profile.get("project_id") or "").strip()

        target = fund_code.strip().upper()
        exact = [item for item in profiles if name(item) == target and project_id(item)]
        if len(exact) == 1:
            return project_id(exact[0])
        if len(exact) > 1:
            raise FundDataError(f"{fund_code}: พบ project ID ซ้ำใน SEC profiles; ระบุ PROJECT ID ในชีต")
        partial = [item for item in profiles
                   if name(item) and project_id(item) and
                   (name(item).startswith(target) or target.startswith(name(item)))]
        if len(partial) == 1:
            return project_id(partial[0])
        if partial:
            raise FundDataError(f"{fund_code}: ชื่อย่อจับคู่ SEC ได้หลายกอง; ระบุ PROJECT ID ในชีต")
        raise FundDataError(f"{fund_code}: ไม่พบใน SEC profiles; ตรวจ AKA หรือ PROJECT ID")

    def nav_items(self, project_id: str, start_date: date, end_date: date) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        cursor: Optional[str] = None
        retry_204 = 0
        page = 0
        while page < MAX_NAV_PAGES:
            params = {
                "proj_id": project_id,
                "start_nav_date": start_date.isoformat(),
                "end_nav_date": end_date.isoformat(),
            }
            if cursor:
                params["next_cursor"] = cursor
            url = f"{SEC_BASE_URL}/daily-info/nav?{urlencode(params)}"
            try:
                status, payload = self._json(url)
            except FundDataError:
                if items:
                    raise
                raise
            if status == 204 and not items and retry_204 < 3:
                retry_204 += 1
                time.sleep(retry_204)
                continue
            if payload is None:
                if items:
                    return items
                raise FundDataError(f"project ID {project_id}: SEC ไม่มีข้อมูล NAV ในช่วงที่ขอ (HTTP {status})")
            page_items = payload.get("items") or payload.get("data") or payload.get("nav_list") or []
            if not isinstance(page_items, list):
                raise FundDataError(f"project ID {project_id}: รูปแบบรายการ NAV จาก SEC ไม่ถูกต้อง")
            items.extend(item for item in page_items if isinstance(item, dict))
            cursor = payload.get("next_cursor") or None
            page += 1
            if not cursor:
                return items
            time.sleep(0.15)
        if cursor:
            raise FundDataError(f"project ID {project_id}: NAV เกินขีดจำกัด {MAX_NAV_PAGES} หน้า")
        return items

    def latest_statistics(self, project_id: str, class_name: str) -> Optional[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        cursor: Optional[str] = None
        for page in range(MAX_FACTSHEET_PAGES):
            params = {"proj_id": project_id, "latest": "true", "page_size": "100"}
            if class_name:
                params["fund_class_name"] = class_name
            if cursor:
                params["next_cursor"] = cursor
            url = f"{SEC_BASE_URL}/factsheet/statistics?{urlencode(params)}"
            status, payload = self._json(url)
            if payload is None:
                if status == 204 and not items:
                    return None
                if status == 204:
                    break
                raise FundDataError(f"project ID {project_id}: SEC ไม่ส่งข้อมูล Factsheet statistics")
            page_items = payload.get("items", [])
            if not isinstance(page_items, list):
                raise FundDataError(f"project ID {project_id}: รูปแบบ Factsheet statistics จาก SEC ไม่ถูกต้อง")
            items.extend(item for item in page_items if isinstance(item, dict))
            cursor = payload.get("next_cursor") or None
            if not cursor:
                break
            time.sleep(0.15)
        else:
            raise FundDataError(
                f"project ID {project_id}: Factsheet statistics เกินขีดจำกัด {MAX_FACTSHEET_PAGES} หน้า"
            )

        if class_name:
            class_items = [
                item for item in items
                if str(item.get("fund_class_name") or "").strip().upper() == class_name.upper()
            ]
        else:
            distinct_classes = {
                str(item.get("fund_class_name") or "").strip().upper()
                for item in items
            }
            if len(distinct_classes) > 1:
                raise FundDataError(
                    f"project ID {project_id}: PTR มีหลาย share class แต่ระบุ class ที่เลือกจาก NAV ไม่ได้"
                )
            class_items = items

        if not class_items:
            return None
        class_items.sort(
            key=lambda item: (
                str(item.get("start_date") or ""),
                str(item.get("last_upd_date") or ""),
            ),
            reverse=True,
        )
        latest = class_items[0]
        if len(class_items) > 1 and (
            str(class_items[1].get("start_date") or "") == str(latest.get("start_date") or "")
            and str(class_items[1].get("last_upd_date") or "")
            == str(latest.get("last_upd_date") or "")
        ):
            raise FundDataError(f"project ID {project_id}: SEC ส่ง PTR ล่าสุดซ้ำในงวดเดียวกัน")
        value = latest.get("portfolio_turnover_ratio")
        if value is None or not str(value).strip():
            return None
        return {
            "value": str(value).strip(),
            "start_date": str(latest.get("start_date") or "").strip(),
        }


def _number(value: Any) -> Optional[float]:
    try:
        result = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _parse_nav_date(value: Any) -> Optional[tuple[str, int]]:
    if value is None:
        return None
    raw = str(value).strip()
    try:
        parsed = date.fromisoformat(raw[:10])
    except ValueError:
        return None
    return parsed.isoformat(), int(datetime(parsed.year, parsed.month, parsed.day,
                                           tzinfo=timezone.utc).timestamp())


def _resolve_class(items: list[dict[str, Any]], hint: str, code: str) -> tuple[list[dict[str, Any]], str]:
    class_field = next((field for field in CLASS_FIELDS
                        if field in items[0] and len({str(item.get(field) or "").strip()
                                                     for item in items}) > 1), None)
    if not class_field:
        return items, ""
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        groups[str(item.get(class_field) or "").strip() or "(unspecified)"].append(item)
    keys = list(groups)
    if hint:
        exact_hint = [key for key in keys if key.upper() == hint.upper()]
        if len(exact_hint) == 1:
            return groups[exact_hint[0]], ""
        matches = [key for key in keys if hint.upper() in key.upper() or key.upper() in hint.upper()]
        if len(matches) == 1:
            return groups[matches[0]], ""
        if len(matches) > 1:
            raise FundDataError(f"{code}: class hint {hint!r} matches multiple SEC classes")
    exact = [key for key in keys if key.upper() == code.upper()]
    if len(exact) == 1:
        return groups[exact[0]], ""
    def non_standard(key: str) -> bool:
        upper = key.upper()
        return any(upper.endswith("-" + suffix) or upper.endswith(suffix)
                   for suffix in NON_STANDARD_CLASS_SUFFIXES)
    standard = [key for key in keys if not non_standard(key)]
    candidates = standard or keys
    largest = max(len(groups[key]) for key in candidates)
    winners = [key for key in candidates if len(groups[key]) == largest]
    if len(winners) != 1:
        raise FundDataError(f"{code}: SEC มีหลาย class ที่จำนวน NAV เท่ากัน; ระบุ class hint ใน PROJECT ID")
    key = winners[0]
    return groups[key], f"เลือก class {key!r} อัตโนมัติจาก {len(keys)} classes"


def get_nav_series(client: SecClient, row: FundRow) -> NavSeries:
    end_date = datetime.now(timezone.utc).date()
    raw_items = client.nav_items(row.project_id, end_date - timedelta(days=HISTORY_DAYS), end_date)
    if not raw_items:
        raise FundDataError(f"{row.code}: SEC ไม่พบข้อมูล NAV")
    selected, warning = _resolve_class(raw_items, row.class_hint, row.code)
    class_field = next((field for field in CLASS_FIELDS if field in selected[0]), None)
    class_name = str(selected[0].get(class_field) or "").strip() if class_field else ""
    by_date: dict[str, tuple[int, float, float]] = {}
    for item in selected:
        parsed = _parse_nav_date(item.get("nav_date") or item.get("navDate") or
                                 item.get("date") or item.get("Date"))
        price = _number(item.get("last_val") or item.get("lastVal") or
                        item.get("nav") or item.get("NAV") or item.get("value"))
        net_asset = _number(item.get("net_asset") or item.get("netAsset") or item.get("NetAsset"))
        if parsed and price is not None and price > 0:
            units = net_asset / price if net_asset is not None else math.nan
            by_date[parsed[0]] = (parsed[1], price, units)
    ordered = [by_date[key] for key in sorted(by_date)]
    if not ordered:
        raise FundDataError(f"{row.code}: ข้อมูล NAV จาก SEC ไม่มีฟิลด์วันที่หรือราคา NAV ที่อ่านได้")
    return NavSeries([item[0] for item in ordered], [item[1] for item in ordered],
                     [item[2] for item in ordered], warning, class_name)


def resample(series: NavSeries, timeframe: str) -> dict[str, Any]:
    seconds, offset = TIMEFRAMES[timeframe]
    buckets: dict[int, dict[str, float]] = {}
    daily_flows = [0.0]
    for index in range(1, len(series.units)):
        current, previous = series.units[index], series.units[index - 1]
        daily_flows.append(abs(current - previous)
                           if math.isfinite(current) and math.isfinite(previous) else 0.0)
    for timestamp, close, flow in zip(series.timestamps, series.closes, daily_flows):
        key = (timestamp + offset) // seconds
        bucket = buckets.setdefault(key, {"high": close, "low": close, "close": close, "volume": 0.0})
        bucket["high"] = max(bucket["high"], close)
        bucket["low"] = min(bucket["low"], close)
        bucket["close"] = close
        bucket["volume"] += flow
    keys = sorted(buckets)
    return {
        "keys": keys,
        "closes": [buckets[key]["close"] for key in keys],
        "highs": [buckets[key]["high"] for key in keys],
        "lows": [buckets[key]["low"] for key in keys],
        "volumes": [buckets[key]["volume"] for key in keys],
    }


def ema(values: list[float], period: int) -> list[Optional[float]]:
    result: list[Optional[float]] = [None] * len(values)
    if len(values) < period:
        return result
    current = sum(values[:period]) / period
    result[period - 1] = current
    multiplier = 2 / (period + 1)
    for index in range(period, len(values)):
        current = (values[index] - current) * multiplier + current
        result[index] = current
    return result


def _last_valid(values: list[Optional[float]], offset: int = 0) -> Optional[float]:
    valid = [value for value in values if value is not None]
    return valid[-1 - offset] if len(valid) > offset else None


def rsi(values: list[float], period: int) -> Optional[float]:
    if len(values) < period + 1:
        return None
    changes = [values[index] - values[index - 1] for index in range(1, len(values))]
    gain = sum(max(value, 0) for value in changes[:period]) / period
    loss = sum(abs(min(value, 0)) for value in changes[:period]) / period
    for change in changes[period:]:
        gain = (gain * (period - 1) + max(change, 0)) / period
        loss = (loss * (period - 1) + abs(min(change, 0))) / period
    return 100.0 if loss == 0 else 100 - 100 / (1 + gain / loss)


def price_structure(highs: list[float], lows: list[float], reference_price: float) -> Optional[dict[str, Any]]:
    def pivots(values: list[float], is_high: bool) -> list[float]:
        found = []
        for index in range(2, len(values) - 2):
            neighbors = values[index - 2:index] + values[index + 1:index + 3]
            if all(values[index] > value if is_high else values[index] < value for value in neighbors):
                found.append(values[index])
        return found

    def compare(swings: list[float], is_high: bool) -> str:
        if len(swings) < 2:
            return ""
        previous, current = swings[-2], swings[-1]
        if previous and abs(current - previous) / abs(previous) <= 0.005:
            return "EQ"
        return ("HH" if current > previous else "LH") if is_high else (
            "HL" if current > previous else "LL")

    swing_highs, swing_lows = pivots(highs, True), pivots(lows, False)
    if len(swing_highs) < 3 or len(swing_lows) < 3:
        return None
    high_state, low_state = compare(swing_highs, True), compare(swing_lows, False)
    previous_high = compare(swing_highs[:-1], True)
    previous_low = compare(swing_lows[:-1], False)
    state = "MIXED"
    if high_state == "HH" and low_state == "HL":
        if previous_high == "HH":
            state = "UP"
        elif previous_high == "LH":
            state = "REVERSAL"
    elif high_state == "LH" and low_state == "LL":
        state = "DOWN"
    elif high_state == "LH" and low_state == "HL":
        if previous_low == "LL":
            state = "DOWN_REV"
        elif previous_high == "HH":
            state = "UP_WEAK"
    if reference_price < swing_lows[-1]:
        state = {"UP": "UP_WEAK", "REVERSAL": "MIXED", "DOWN_REV": "DOWN",
                 "UP_WEAK": "DOWN"}.get(state, state)
    return {"state": state, "high": swing_highs[-1], "low": swing_lows[-1]}


def analyze_timeframe(series: dict[str, Any], timeframe: str, reference_price: float) -> dict[str, Any]:
    values = series["closes"]
    config = TF_CONFIG[timeframe]
    fast = ema(values, config["fast"])
    slow = ema(values, config["slow"])
    macd_fast = ema(values, config["macd_fast"])
    macd_slow = ema(values, config["macd_slow"])
    macd = [a - b if a is not None and b is not None else None
            for a, b in zip(macd_fast, macd_slow)]
    first_valid = next((index for index, value in enumerate(macd) if value is not None), len(macd))
    compact_macd = [value for value in macd[first_valid:] if value is not None]
    signal_compact = ema(compact_macd, config["macd_signal"])
    signal = [None] * first_valid + signal_compact
    signal += [None] * (len(macd) - len(signal))
    histogram = [a - b if a is not None and b is not None else None
                 for a, b in zip(macd, signal)]
    last_hist, previous_hist = _last_valid(histogram), _last_valid(histogram, 1)
    delta = last_hist - previous_hist if last_hist is not None and previous_hist is not None else None
    trend = None
    if len(values) >= config["trend"] + 1:
        trend = "UP" if values[-1] >= sum(values[-config["trend"]:]) / config["trend"] else "DOWN"
    current_rsi = rsi(values, config["rsi"])
    cross = ""
    if len(fast) > 1 and all(value is not None for value in (fast[-1], slow[-1], fast[-2], slow[-2])):
        if fast[-2] <= slow[-2] and fast[-1] > slow[-1]:
            cross = "BUY"
        elif fast[-2] >= slow[-2] and fast[-1] < slow[-1]:
            cross = "SELL"
    macd_cross = ""
    if len(macd) > 1 and all(value is not None for value in
                             (macd[-1], signal[-1], macd[-2], signal[-2])):
        if macd[-2] <= signal[-2] and macd[-1] > signal[-1]:
            macd_cross = "BUY"
        elif macd[-2] >= signal[-2] and macd[-1] < signal[-1]:
            macd_cross = "SELL"
    sig_rsi = ""
    if trend == "UP" and current_rsi is not None and current_rsi < config["buy_rsi"]:
        sig_rsi = "BUY"
    elif trend == "DOWN" and current_rsi is not None and current_rsi > config["sell_rsi"]:
        sig_rsi = "SELL"
    hist_improved = (len(histogram) >= 2 and histogram[-1] is not None and
                     histogram[-2] is not None and histogram[-1] > histogram[-2])
    hist_worsened = (len(histogram) >= 2 and histogram[-1] is not None and
                     histogram[-2] is not None and histogram[-1] < histogram[-2])
    sig_macd = ""
    if (current_rsi is not None and current_rsi < config["buy_rsi"] and hist_improved and
            macd_cross == "BUY" and cross == "BUY"):
        sig_macd = "BUY"
    elif (current_rsi is not None and current_rsi > config["sell_rsi"] and hist_worsened and
          macd_cross == "SELL" and cross == "SELL"):
        sig_macd = "SELL"
    sig_volume = None
    volumes = series["volumes"]
    ref_index = len(volumes) - 2
    if len(volumes) >= config["vol"] + 2 and ref_index - config["vol"] >= 0:
        average = sum(volumes[ref_index - config["vol"]:ref_index]) / config["vol"]
        if average:
            if volumes[ref_index] > average and volumes[ref_index] > volumes[ref_index - 1]:
                sig_volume = "BUY"
            elif (values[ref_index] < values[ref_index - 1] and delta is not None and
                  delta < 0 and cross == "SELL" and volumes[ref_index] < average * 0.8):
                sig_volume = "SELL"
            else:
                sig_volume = "HOLD"
    structure = price_structure(series["highs"], series["lows"], reference_price)
    return {
        "trend": trend, "rsi": current_rsi, "ema_fast": _last_valid(fast),
        "ema_slow": _last_valid(slow), "hist_delta": delta, "sig_volume": sig_volume,
        "sig_rsi": sig_rsi, "sig_macd": sig_macd, "sig_ema": cross,
        "structure": structure, "last_close": values[-1],
        "macd": _last_valid(macd), "macd_signal": _last_valid(signal),
        "macd_hist": last_hist,
    }


def build_report(row: FundRow, nav: NavSeries) -> dict[str, Any]:
    series = {timeframe: resample(nav, timeframe) for timeframe in TIMEFRAMES}
    results = {}
    finer_timeframe = {"W": "D", "D": "4H", "4H": "1H"}
    for timeframe, tf_series in series.items():
        finer = series.get(finer_timeframe.get(timeframe, ""))
        reference_series = finer if finer and len(finer["closes"]) >= 2 else tf_series
        reference = (reference_series["closes"][-2] if len(reference_series["closes"]) >= 2
                     else reference_series["closes"][-1])
        results[timeframe] = analyze_timeframe(tf_series, timeframe, reference)
    weekly_structure = results["W"]["structure"]
    bias = ""
    if weekly_structure:
        bias = {"UP": "BULLISH", "DOWN": "BEARISH", "REVERSAL": "RECOVERY",
                "DOWN_REV": "RECOVERY"}.get(weekly_structure["state"], "NEUTRAL")
    daily = results["D"]
    setup = ""
    structure = daily["structure"]
    if structure:
        if daily["rsi"] is not None and daily["rsi"] > TF_CONFIG["D"]["sell_rsi"]:
            setup = "EXTENDED"
        elif structure["state"] == "DOWN":
            setup = "FAILED"
        elif structure["state"] == "DOWN_REV":
            setup = "RECOVERY"
        elif structure["state"] in ("UP", "REVERSAL"):
            setup = "PULLBACK" if daily["rsi"] is not None and daily["rsi"] <= 50 else "READY"
        else:
            setup = "WAIT"
    zone = ""
    structure4h = results["4H"]["structure"]
    if structure4h:
        state = structure4h["state"]
        close = results["4H"]["last_close"]
        if state == "UP":
            zone = "🟢 BREAKOUT" if close > structure4h["high"] else "🟢 HH/HL"
        elif state in ("REVERSAL", "DOWN_REV"):
            zone = "🟡 RETEST"
        elif state == "UP_WEAK":
            zone = "🟡 RANGE"
        elif state == "DOWN":
            zone = "🔴 BREAKDOWN" if close < structure4h["low"] else "🔴 LH/LL"
        else:
            zone = "🟡 RANGE"
    trigger = ""
    structure1h = results["1H"]["structure"]
    one_hour = results["1H"]
    if structure1h:
        state = structure1h["state"]
        bull = one_hour["sig_ema"] == "BUY" and (
            one_hour["sig_volume"] == "BUY" or
            (one_hour["hist_delta"] is not None and one_hour["hist_delta"] > 0))
        bear = one_hour["sig_ema"] == "SELL" and (
            one_hour["sig_volume"] == "SELL" or
            (one_hour["hist_delta"] is not None and one_hour["hist_delta"] < 0))
        if state in ("UP", "REVERSAL"):
            trigger = ("🟢 BREAKOUT" if one_hour["last_close"] > structure1h["high"] else "🟢 ENTRY") if bull else "🟡 WAIT"
        elif state in ("DOWN", "UP_WEAK"):
            trigger = "🔴 FAILED" if bear else "🟡 RETEST"
        else:
            trigger = "🟡 WAIT"
    action = ""
    zone_bull, zone_bear = zone.startswith("🟢"), zone.startswith("🔴")
    trigger_bull, trigger_bear = trigger.startswith("🟢"), trigger.startswith("🔴")
    trigger_wait = trigger.startswith("🟡")
    if bias == "BULLISH" and setup in ("READY", "RECOVERY", "PULLBACK"):
        if zone_bull and trigger_bull:
            action = "BUY"
        elif zone_bull and trigger_wait:
            action = "WAIT"
    elif bias == "BEARISH" and setup == "FAILED":
        if zone_bear and trigger_bear:
            action = "SELL"
        elif zone_bear:
            action = "WAIT"
    return {
        "row": row, "nav": nav, "series": series, "timeframes": results,
        "bias": bias, "setup": setup, "zone": zone, "trigger": trigger,
        "action": action,
    }


def _format_number(value: Any, decimals: int = 4) -> str:
    if value is None:
        return "—"
    return f"{value:.{decimals}f}"


def _format_ptr(report: dict[str, Any]) -> str:
    ptr = report.get("ptr")
    if not ptr:
        return "N/A"
    value = str(ptr["value"]).replace("|", "\\|")
    start_date = str(ptr.get("start_date") or "").strip()
    return f"{value} ({start_date})" if start_date else value


def _fund_report_row(report: dict[str, Any], specified: bool = False) -> tuple[str, ...]:
    metrics = report["timeframes"]
    row = report["row"]
    nav = report["nav"]
    code = row.code
    comparison_timeframe = "1H" if specified else "D"
    return (
        f"[{code}]({FINNOMENA_URL.format(quote(code, safe=''))})",
        row.project_id or "—",
        _format_number(metrics["W"]["rsi"], 2),
        _format_number(metrics[comparison_timeframe]["rsi"], 2),
        _format_number(metrics["W"]["hist_delta"], 5),
        _format_number(nav.closes[-1], 6),
        report["bias"] or "WAIT",
        report["setup"] or "WAIT",
        report["zone"] or "WAIT",
        report["trigger"] or "WAIT",
        {"SELL": "EXIT"}.get(report["action"], report["action"]) or "—",
        _format_ptr(report),
    )


def _print_fund_metadata(reports: list[dict[str, Any]]) -> None:
    updated = datetime.now(BANGKOK_TIMEZONE).strftime("%d/%m/%Y %H:%M")
    for index, report in enumerate(reports):
        nav = report["nav"]
        candle_date = datetime.fromtimestamp(
            nav.timestamps[-1], BANGKOK_TIMEZONE
        ).strftime("%d/%m/%Y")
        print(f"AKA         : {report['row'].code}")
        print(f"LAST_CANDLE : {candle_date} (NAV รายวัน)")
        print(f"UPDATED     : {updated}")
        print(f"PROJECT ID  : {report['row'].project_id or '—'}")
        if index < len(reports) - 1:
            print()


def print_update(reports: list[dict[str, Any]], specified: bool = False) -> None:
    if specified:
        _print_fund_metadata(reports)
        print(f"\n🔵 SPECIFIED FUND — {len(reports)} รายการ")
    else:
        print(f"\nสรุปข้อมูลกองทุน — {len(reports)} กองทุน")
    comparison_rsi_header = "RSI(1H)" if specified else "RSI(D)"
    headers = (
        "AKA", "PROJECT ID", "RSI(W)", comparison_rsi_header, "MACD Δ(W)", "NAV",
        "BIAS(W)", "SETUP(D)", "ENTRY_ZONE(4H)", "ENTRY_TRIGGER(1H)",
        "ACTION", "PTR (latest Factsheet)",
    )
    print("| " + " | ".join(headers) + " |")
    print("|" + "|".join("---" for _ in headers) + "|")
    for report in reports:
        print("| " + " | ".join(_fund_report_row(report, specified)) + " |")
    warnings = [report["nav"].class_warning for report in reports if report["nav"].class_warning]
    if warnings:
        print("\nSEC class warnings:")
        for report in reports:
            if report["nav"].class_warning:
                print(f"- {report['row'].code}: {report['nav'].class_warning}")
    print(f"\nUPDATE — {len(reports)} กองทุน | ไม่ได้เขียนข้อมูลกลับไปยังชีต")


def candidate_report_groups(reports: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    primed, near_zone = [], []
    for report in reports:
        if report["bias"] != "BULLISH" or report["setup"] not in ("READY", "RECOVERY", "PULLBACK"):
            continue
        if report["zone"].startswith("🟢") and report["trigger"].startswith("🟡"):
            primed.append(report)
        elif not report["zone"].startswith("🟢") and not report["trigger"].startswith("🟢"):
            near_zone.append(report)
    return primed, near_zone


def reports_requiring_ptr(command: str, reports: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if command == "UPDATE":
        return reports
    return [report for group in candidate_report_groups(reports) for report in group]


def print_candidates(reports: list[dict[str, Any]]) -> None:
    primed, near_zone = candidate_report_groups(reports)
    for title, selected in (
        ("🟢 พร้อมครบ W→D→4H — รอ ENTRY_TRIGGER(1H)", primed),
        ("🟡 W→D ผ่านแล้ว แต่ 4H ยังไม่ยืนยันโซน", near_zone),
    ):
        print(f"\n{title} — {len(selected)} กองทุน")
        headers = (
            "#", "AKA", "PROJECT ID", "NAV", "MACD Δ(W)",
            "ENTRY_ZONE(4H)", "ENTRY_TRIGGER(1H)", "PTR (latest Factsheet)",
        )
        print("| " + " | ".join(headers) + " |")
        print("|" + "|".join("---" for _ in headers) + "|")
        if not selected:
            print("| " + " | ".join("—" for _ in headers) + " |")
            continue
        ranked = sorted(
            selected,
            key=lambda report: (
                report["timeframes"]["W"]["hist_delta"] is None,
                -(report["timeframes"]["W"]["hist_delta"] or 0),
            ),
        )
        for rank, report in enumerate(ranked, 1):
            code = report["row"].code
            values = (
                str(rank),
                f"[{code}]({FINNOMENA_URL.format(quote(code, safe=''))})",
                report["row"].project_id or "—",
                _format_number(report["nav"].closes[-1], 6),
                _format_number(report["timeframes"]["W"]["hist_delta"], 5),
                report["zone"] or "WAIT",
                report["trigger"] or "WAIT",
                _format_ptr(report),
            )
            print("| " + " | ".join(values) + " |")


def _parse_command(command: str) -> tuple[str, Optional[list[str]]]:
    if command == "UPDATE" or command == "CANDIDATE":
        return command, None
    if command.startswith("UPDATE/"):
        symbols = command[len("UPDATE/"):].split(",")
        if not symbols or any(not FUND_CODE_RE.fullmatch(symbol) for symbol in symbols):
            raise FundDataError("รูปแบบคำสั่งไม่ถูกต้อง: UPDATE/<AKA>,<AKA>,... (ใช้ comma คั่น)")
        if len({symbol.upper() for symbol in symbols}) != len(symbols):
            raise FundDataError("มี AKA ซ้ำในคำสั่ง UPDATE")
        return "UPDATE", symbols
    raise FundDataError("คำสั่งต้องเป็น UPDATE, UPDATE/<AKA>,... หรือ CANDIDATE")


def main(arguments: list[str]) -> int:
    if len(arguments) > 1:
        print("Usage: fund_update.py [UPDATE|UPDATE/<AKA>,...|CANDIDATE]", file=sys.stderr)
        return 2
    try:
        command, selected = _parse_command(arguments[0] if arguments else "UPDATE")
        config = load_config()
        rows = load_fund_rows(config, selected)
        print(f"✅ ดาวน์โหลดและตรวจสอบ snapshot: {WORKBOOK_PATH}")
        client = SecClient(config["SEC_API_KEY_FUND"])
        if any(not row.project_id for row in rows):
            print("กำลังค้นหา PROJECT ID ที่ว่างจาก SEC profiles...", file=sys.stderr)
            profiles = client.profiles()
            rows = [FundRow(row.code, row.project_id or client.find_project_id(profiles, row.code),
                            row.class_hint, row.row_number) for row in rows]
        reports: list[dict[str, Any]] = []
        errors: list[str] = []
        for index, row in enumerate(rows, 1):
            print(f"กำลังดึง NAV {index}/{len(rows)}: {row.code}", file=sys.stderr)
            try:
                reports.append(build_report(row, get_nav_series(client, row)))
            except FundDataError as exc:
                errors.append(str(exc))
            if index < len(rows):
                time.sleep(0.25)
        ptr_reports = reports_requiring_ptr(command, reports)
        for index, report in enumerate(ptr_reports, 1):
            row = report["row"]
            class_name = report["nav"].class_name or row.class_hint
            print(f"กำลังดึง PTR {index}/{len(ptr_reports)}: {row.code}", file=sys.stderr)
            try:
                report["ptr"] = client.latest_statistics(row.project_id, class_name)
            except FundDataError as exc:
                report["ptr"] = None
                errors.append(f"{row.code}: ดึง PTR จาก SEC ไม่สำเร็จ: {exc}")
            if index < len(ptr_reports):
                time.sleep(0.15)
        if reports:
            print_update(reports, specified=selected is not None) if command == "UPDATE" else print_candidates(reports)
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        if errors:
            print(
                f"สร้างรายงานให้ {len(reports)} กองทุน; พบข้อผิดพลาด {len(errors)} รายการ",
                file=sys.stderr,
            )
            return 1
        if command == "CANDIDATE":
            print(f"\nตรวจแล้ว {len(reports)} กองทุน")
        return 0
    except FundDataError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
